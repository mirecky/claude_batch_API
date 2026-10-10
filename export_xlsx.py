"""export_xlsx.py – prevedie results_*.jsonl do Excelu.

Použitie:
  py export_xlsx.py                          # všetky results_*.jsonl v zložke → vysledky.xlsx
  py export_xlsx.py results_msgbatch_X.jsonl --out dotazniky.xlsx
  py export_xlsx.py --rows entries           # 1 riadok = 1 prvok poľa 'entries' (ak ho schéma má)

Hárok 'data'   : 1 riadok = 1 strana (alebo 1 prvok poľa --rows); vnorené polia ako JSON text.
Hárok 'kontrola': strany s chybou / upozornením / bez dát.
Zoradené podľa custom_id (= poradie strán v PDF). Ak sú v pages/manifest.csv, doplní pdf a stranu.
Pre konkrétnu aplikáciu (presné stĺpce) si dajte v chate upraviť funkciu row_for().
"""
import argparse, csv, json, sys
from pathlib import Path
try:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("Chýba balík. Spusti:  py -m pip install openpyxl")

def flat(v):
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False)
    return v

def row_for(base, d):
    r = dict(base)
    for k, v in (d or {}).items():
        r[k] = flat(v)
    return r

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("files", nargs="*")
    ap.add_argument("--out", default="vysledky.xlsx")
    ap.add_argument("--rows", help="názov poľa, ktoré sa rozbalí na viac riadkov (napr. entries)")
    a = ap.parse_args()
    files = [Path(f) for f in a.files] or sorted(Path(".").glob("results_*.jsonl"))
    if not files:
        sys.exit("Nenašiel som results_*.jsonl")
    man = {}
    mf = Path("pages/manifest.csv")
    if mf.exists():
        with mf.open(encoding="utf-8") as f:
            man = {r["id"]: r for r in csv.DictReader(f, delimiter=";")}
    rows, check = [], []
    seen = {}
    for fp in files:                                  # neskorší súbor prepíše skorší (napr. opakované čítanie)
        for line in fp.open(encoding="utf-8"):
            r = json.loads(line); seen[r["custom_id"]] = r
    for cid in sorted(seen):
        r = seen[cid]
        base = {"id": cid, "pdf": man.get(cid, {}).get("pdf", ""), "strana": man.get(cid, {}).get("strana", "")}
        d = r.get("data")
        if r.get("error") or not d or r.get("note"):
            check.append({**base, "problem": r.get("error") or r.get("note") or "bez dát"})
        if not d:
            continue
        if a.rows and isinstance(d.get(a.rows), list) and d[a.rows]:
            top = {k: v for k, v in d.items() if k != a.rows}
            for i, e in enumerate(d[a.rows], 1):
                rows.append(row_for({**base, "poradie": i}, {**top, **(e if isinstance(e, dict) else {"hodnota": e})}))
        else:
            rows.append(row_for(base, d))
    wb = Workbook()
    for ws, data in ((wb.active, rows), (wb.create_sheet("kontrola"), check)):
        ws.title = ws.title if ws.title == "kontrola" else "data"
        cols = []
        for r in data:
            cols += [k for k in r if k not in cols]
        ws.append(cols or ["(prázdne)"])
        for r in data:
            ws.append([r.get(c, "") for c in cols])
        ws.freeze_panes = "A2"
        for i, c in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, len(str(c)) + 2))
    wb.save(a.out)
    print(f"{a.out}: {len(rows)} riadkov, na kontrolu {len(check)}")

if __name__ == "__main__":
    main()
